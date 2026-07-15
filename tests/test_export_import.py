from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.extensions import db
from app.models import Flat

from .factories import make_flat, make_job


def test_export_excludes_soft_deleted_jobs(client):
    make_job(title="Kept job")
    dead = make_job(title="Binned job")
    dead.soft_delete()
    db.session.commit()

    response = client.get("/jobs/export")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["Content-Type"]

    workbook = load_workbook(BytesIO(response.data))
    titles = [row[5] for row in workbook["All Jobs"].iter_rows(min_row=2, values_only=True)]
    assert "Kept job" in titles
    assert "Binned job" not in titles


def _import_workbook(rows: list[dict]) -> BytesIO:
    headers = [
        "Payment Reference",
        "Flat No.",
        "Tenancy / Property Address",
        "Tenancy / Property Postcode",
        "Tenants",
        "Tenant Contact Details",
        "Rent",
        "Due Date",
        "Landlord",
        "Tenancy Start Date",
        "Tenancy End Date",
    ]
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def test_import_creates_and_updates_by_payment_reference(client):
    make_flat(payment_reference="EXISTING-1", address="Old Address")

    file = _import_workbook(
        [
            {
                "Payment Reference": "EXISTING-1",
                "Flat No.": "1A",
                "Tenancy / Property Address": "Updated Address",
                "Rent": "£1,300.00",
                "Due Date": "15th",
                "Tenancy Start Date": "01/02/2024",
            },
            {
                "Payment Reference": "BRAND-NEW-1",
                "Flat No.": "2B",
                "Tenancy / Property Address": "New Address",
                "Rent": "950",
            },
            {"Payment Reference": "", "Tenancy / Property Address": ""},  # skipped
        ]
    )

    response = client.post(
        "/flats/import",
        data={"file": (file, "flats.xlsx")},
        content_type="multipart/form-data",
    )
    assert response.status_code == 302

    existing = db.session.scalar(db.select(Flat).where(Flat.payment_reference == "EXISTING-1"))
    assert existing.address == "Updated Address"
    assert existing.rent_amount == 1300.0
    assert existing.rent_due_date == 15
    assert str(existing.tenancy_start_date) == "2024-02-01"

    new = db.session.scalar(db.select(Flat).where(Flat.payment_reference == "BRAND-NEW-1"))
    assert new is not None
    assert new.rent_amount == 950.0
